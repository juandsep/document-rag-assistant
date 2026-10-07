"""Unit tests for loading, cleaning and chunking each supported format."""

import pytest

from rag.ingest import chunk_document, chunk_text, clean, load_pages


def _pdf(path, pages):
    """Write a minimal valid PDF with one line of Helvetica text per page."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects)} 0 R >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))


def test_chunk_text_respects_size_and_overlap():
    text = "a" * 1000
    chunks = chunk_text(text, doc_id="d1", size=400, overlap=50)

    assert [len(c.text) for c in chunks] == [400, 400, 300]
    assert [c.index for c in chunks] == [0, 1, 2]
    assert chunks[0].text == text[:400]


def test_chunk_text_never_emits_a_chunk_that_only_repeats_the_previous_tail():
    # 1000 chars, step 300: a fourth window at 900 would sit inside the third.
    assert len(chunk_text("a" * 1000, size=400, overlap=100)) == 3


def test_chunk_text_rejects_bad_arguments():
    with pytest.raises(ValueError):
        chunk_text("abc", size=0)
    with pytest.raises(ValueError):
        chunk_text("abc", size=10, overlap=10)


def test_chunk_text_of_empty_text_is_empty():
    assert chunk_text("") == []


def test_clean_rejoins_hyphenated_words_and_collapses_whitespace():
    assert clean("devolu-\nciones   se\n\naceptan ") == "devoluciones se aceptan"


def test_markdown_loads_as_one_page(tmp_path):
    (tmp_path / "faq.md").write_text("# FAQ\n\nReturns: 30 days.", encoding="utf-8")
    assert load_pages(tmp_path / "faq.md") == [(None, "# FAQ\n\nReturns: 30 days.")]


def test_pdf_keeps_page_numbers_and_numbers_chunks_continuously(tmp_path):
    _pdf(
        tmp_path / "policy.pdf",
        ["Returns are accepted for 30 days.", "Shipping takes 5 days."],
    )

    chunks = chunk_document(tmp_path / "policy.pdf", "policy.pdf")

    assert [(c.page, c.index) for c in chunks] == [(1, 0), (2, 1)]
    assert "30 days" in chunks[0].text and "5 days" in chunks[1].text


def test_docx_reads_paragraphs_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Warranty: 12 months.")
    document.add_table(rows=1, cols=2).rows[0].cells[1].text = "Furniture: 24 months"
    document.save(tmp_path / "warranty.docx")

    [chunk] = chunk_document(tmp_path / "warranty.docx", "warranty.docx")

    assert "12 months" in chunk.text and "24 months" in chunk.text
    assert chunk.page is None


def test_unsupported_formats_fail_loudly(tmp_path):
    (tmp_path / "data.csv").write_text("a,b")
    with pytest.raises(ValueError, match="Unsupported"):
        load_pages(tmp_path / "data.csv")


def test_with_translations_adds_a_copy_searchable_in_the_other_language():
    from rag.ingest import Chunk, with_translations

    chunks = [
        Chunk("es.txt", 0, "La garantía dura 12 meses."),
        Chunk("en.txt", 0, "Shipping takes 5 days."),
    ]
    calls = []

    def translate(text, target):
        calls.append(target)
        return f"<{target}> {text}"

    out = with_translations(chunks, translate)

    assert calls == ["English", "Spanish"]
    assert out[:2] == chunks
    assert [(c.doc_id, c.text, c.search_text) for c in out[2:]] == [
        (
            "es.txt",
            "La garantía dura 12 meses.",
            "<English> La garantía dura 12 meses.",
        ),
        ("en.txt", "Shipping takes 5 days.", "<Spanish> Shipping takes 5 days."),
    ]
