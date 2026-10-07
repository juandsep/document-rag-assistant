"""Document ingestion: load, clean and chunk.

Plain text, Markdown, digital PDF and DOCX load into pages of text (a PDF keeps
its page numbers, so answers can cite them). Scanned PDFs are not read: there
is no OCR, and `textless_pages` lets indexing report them. The PDF and DOCX parsers live in the
`ingest` dependency group: indexing needs them, the API image does not.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, replace
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


@dataclass(frozen=True)
class Chunk:
    doc_id: str
    index: int
    text: str
    page: int | None = None
    # What gets embedded when it differs from `text` (a translation of it);
    # the stored, cited text is always the original.
    search_text: str | None = None


# Spanish markers: inverted punctuation, accents, frequent function words.
_SPANISH = re.compile(
    r"[¿¡áéíóúñ]|\b(el|la|los|las|de|que|qué|cómo|cuánto|cuántos|cuál|puedo|"
    r"tengo|es|un|una|mi|si|se|por|para|con|hay)\b",
    re.IGNORECASE,
)


def language(text: str) -> str:
    """Spanish or English: the two languages the corpus and the UI support."""
    return "Spanish" if _SPANISH.search(text) else "English"


def with_translations(
    chunks: list[Chunk], translate: Callable[[str, str], str]
) -> list[Chunk]:
    """Add, for each chunk, a copy searchable in the other language.

    The multilingual embedding ranks a Spanish question about an English
    document (and the reverse) below same-language passages; on the eval set
    every retrieval miss was cross-language. Embedding a translation of each
    chunk next to the original fixes that at index time, with no cost per
    query: recall@3 rose from 0.862 to 0.968.
    """
    translated = []
    for chunk in chunks:
        target = "English" if language(chunk.text) == "Spanish" else "Spanish"
        translated.append(replace(chunk, search_text=translate(chunk.text, target)))
    return chunks + translated


def clean(text: str) -> str:
    """Undo PDF line wrapping: rejoin hyphenated words, collapse whitespace."""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    return re.sub(r"\s+", " ", text).strip()


def chunk_text(
    text: str,
    doc_id: str = "doc",
    size: int = 800,
    overlap: int = 100,
    page: int | None = None,
    start_index: int = 0,
) -> list[Chunk]:
    """Split `text` into `size`-character chunks that overlap by `overlap`.

    The overlap keeps a fact that straddles a boundary whole in one chunk.
    Indexes start at `start_index`, so a multi-page document numbers its
    chunks continuously.
    """
    if size <= 0:
        raise ValueError("size must be > 0")
    if not 0 <= overlap < size:
        raise ValueError("overlap must satisfy 0 <= overlap < size")

    step = size - overlap
    chunks: list[Chunk] = []
    for start in range(0, len(text), step):
        chunks.append(
            Chunk(doc_id, start_index + len(chunks), text[start : start + size], page)
        )
        # The next window would only repeat this one's tail.
        if start + size >= len(text):
            break
    return chunks


def load_pages(path: Path) -> list[tuple[int | None, str]]:
    """Return `(page, text)` pairs; page is None where the format has none."""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [(None, path.read_text(encoding="utf-8"))]
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return [
            (n, page.extract_text() or "") for n, page in enumerate(reader.pages, 1)
        ]
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return [(None, "\n".join([p.text for p in document.paragraphs] + cells))]
    raise ValueError(
        f"Unsupported format {suffix!r}; expected one of {sorted(SUPPORTED)}"
    )


def textless_pages(pages: list[tuple[int | None, str]]) -> list[int]:
    """Pages with no extractable text: scanned images, in a PDF.

    Only digital PDFs (with a text layer) are supported; there is no OCR, so
    indexing reports these instead of silently storing nothing.
    """
    return [page for page, text in pages if page is not None and not clean(text)]


def chunk_pages(pages: list[tuple[int | None, str]], doc_id: str) -> list[Chunk]:
    """Clean and chunk loaded pages; empty pages yield nothing."""
    chunks: list[Chunk] = []
    for page, text in pages:
        chunks += chunk_text(clean(text), doc_id, page=page, start_index=len(chunks))
    return chunks


def chunk_document(path: Path, doc_id: str) -> list[Chunk]:
    """Load, clean and chunk one file."""
    return chunk_pages(load_pages(path), doc_id)
